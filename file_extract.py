#!/usr/bin/env python3
"""
file_extract.py - extract text from NON media (image/video) attachments so Hadi
can "see" documents sent in the channel: PDF, Word, Excel, and text-like files.

Images and videos are handled elsewhere (save_image_attachments / save_video_frames).
This module covers everything else. Fail-open and lazy imports: a missing library
degrades to a short note instead of crashing the bot.

Usage (from discord_bot on_message, AFTER the speak-gate so we don't download
files for messages we skip):

    import file_extract
    doc_notes = await file_extract.extract_attachment_texts(message)
    if doc_notes:
        media_notes = media_notes + "\\n" + "\\n".join(doc_notes)
"""
import io

TEXT_EXTS = {
    "txt", "csv", "tsv", "json", "md", "log", "py", "js", "ts", "tsx", "jsx",
    "html", "htm", "xml", "yml", "yaml", "sql", "ini", "cfg", "conf", "sh",
    "java", "kt", "cs", "go", "rb", "php", "c", "cpp", "h", "env",
}
MAX_FILES = 6
MAX_CHARS = 6000


def _clip(s, n=MAX_CHARS):
    s = s or ""
    return s if len(s) <= n else s[:n] + "\n…(اتقصّ)"


def _pdf(data):
    try:
        import pypdf
        r = pypdf.PdfReader(io.BytesIO(data))
        return "\n".join((p.extract_text() or "") for p in r.pages[:30]).strip()
    except Exception as e:  # noqa
        return f"(PDF - مش قادر أقراه: {e})"


def _docx(data):
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in d.paragraphs if p.text).strip()
    except Exception as e:  # noqa
        return f"(Word - مش قادر أقراه: {e})"


def _xlsx(data):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:5]:
            out.append(f"# ورقة: {ws.title}")
            for r in ws.iter_rows(max_row=200, values_only=True):
                cells = [str(c) for c in r if c is not None]
                if cells:
                    out.append(" | ".join(cells))
        return "\n".join(out).strip()
    except Exception as e:  # noqa
        return f"(Excel - مش قادر أقراه: {e})"


def _text(data):
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(enc).strip()
        except Exception:  # noqa
            continue
    return "(نص - مش قادر أفك ترميزه)"


async def extract_attachment_texts(message, max_files=MAX_FILES):
    notes = []
    atts = getattr(message, "attachments", None) or []
    for att in atts[:max_files]:
        ct = (getattr(att, "content_type", "") or "").lower()
        name = getattr(att, "filename", "file") or "file"
        ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
        if ct.startswith("image/") or ct.startswith("video/"):
            continue  # handled elsewhere
        try:
            data = await att.read()
        except Exception as e:  # noqa
            notes.append(f"[ملف {name}] مش قادر أنزّله: {e}")
            continue
        if ext == "pdf" or ct == "application/pdf":
            body = _pdf(data)
        elif ext == "docx" or "word" in ct:
            body = _docx(data)
        elif ext in ("xlsx", "xlsm") or "sheet" in ct or "excel" in ct:
            body = _xlsx(data)
        elif ext in TEXT_EXTS or ct.startswith("text/") or ct in ("application/json",):
            body = _text(data)
        else:
            size = getattr(att, "size", 0)
            notes.append(f"[ملف {name}] نوعه ({ct or ext or '؟'})، {size} بايت — مش نوع أقدر أقرا محتواه نصيًا.")
            continue
        notes.append(f"[محتوى ملف {name}]:\n{_clip(body)}")
    return notes
